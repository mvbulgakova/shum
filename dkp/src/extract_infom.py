"""Извлекает временные ряды инфляционных ожиданий инФОМ из XLSX.

Источник: https://www.cbr.ru/analytics/dkp/inflationary_expectations/ (латест-файл)
Лист «Данные за все годы» содержит все месячные наблюдения с 2009 года.

Выход: data/processed/infom.csv с колонками
  month           — YYYY-MM-01
  observed_med    — медиана наблюдаемой инфляции, %
  expected_med    — медиана ожидаемой инфляции (на 12 мес), %
  expected_5y_med — медиана ожидаемой пятилетней инфляции, %
"""

from __future__ import annotations

import csv
import datetime
import sys
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw" / "infom" / "latest.xlsx"
OUT = ROOT / "data" / "processed" / "infom.csv"
URL = "https://www.cbr.ru/analytics/dkp/inflationary_expectations/"


def ensure_download() -> None:
    """Скачивает самый свежий xlsx (парсит листинг)."""
    import re
    RAW.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=25) as r:
        html = r.read().decode("utf-8", "replace")
    m = re.search(r'href="(/Collection/Collection/File/\d+/Infl_exp_\d\d-\d\d\.xlsx)"', html)
    if not m:
        raise RuntimeError("Не нашёл ссылку на латест-файл на " + URL)
    xlsx_url = "https://www.cbr.ru" + m.group(1)
    print(f"Скачиваю: {xlsx_url}")
    urllib.request.urlretrieve(xlsx_url, RAW)


def main() -> int:
    if not RAW.exists() or RAW.stat().st_size < 100_000:
        ensure_download()

    wb = openpyxl.load_workbook(RAW, data_only=True)
    ws = wb["Данные за все годы"]

    # Row 2 содержит даты (в формате datetime) по столбцам.
    dates_row = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    date_cols: list[tuple[int, datetime.date]] = [
        (i, v.date()) for i, v in enumerate(dates_row) if isinstance(v, datetime.datetime)
    ]

    # Нам нужны три ряда:
    #   1) блок «Прямые оценки годовой инфляции: медианные значения»
    #      → «наблюдаемая инфляция (в %)» и «ожидаемая инфляция (в %)»
    #   2) блок «Прямые оценки пятилетней инфляции: медианные значения»
    #      → «ожидаемая инфляция (в %)»
    series: dict[str, list[float | None]] = {
        "observed_med": [None] * len(date_cols),
        "expected_med": [None] * len(date_cols),
        "expected_5y_med": [None] * len(date_cols),
    }

    current_section: str = ""
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        label = row[0]
        if label is None:
            continue
        ls = str(label).strip()

        # Секции — жирные заголовки с фразой "медианные значения"
        if "Прямые оценки годовой инфляции" in ls:
            current_section = "annual"
            continue
        if "Прямые оценки пятилетней инфляции" in ls:
            current_section = "5y"
            continue
        # выход из релевантных секций
        if "Прямые оценки" in ls or "Медианное значение" in ls:
            current_section = ""

        low = ls.lower()
        key: str | None = None
        if current_section == "annual" and low.startswith("наблюдаемая инфляция"):
            key = "observed_med"
        elif current_section == "annual" and low.startswith("ожидаемая инфляция"):
            key = "expected_med"
        elif current_section == "5y" and low.startswith("ожидаемая инфляция"):
            key = "expected_5y_med"
        if key is None:
            continue
        # разбираем значения по столбцам-датам
        for i, (col, _) in enumerate(date_cols):
            v = row[col] if col < len(row) else None
            if isinstance(v, (int, float)):
                series[key][i] = round(float(v), 3)

    with OUT.open("w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["month", "observed_med", "expected_med", "expected_5y_med"])
        for i, (_, d) in enumerate(date_cols):
            w.writerow([
                d.isoformat(),
                series["observed_med"][i] if series["observed_med"][i] is not None else "",
                series["expected_med"][i] if series["expected_med"][i] is not None else "",
                series["expected_5y_med"][i] if series["expected_5y_med"][i] is not None else "",
            ])

    filled = sum(1 for v in series["expected_med"] if v is not None)
    print(f"OK: {len(date_cols)} месяцев, ожидаемая заполнена в {filled} из них.")
    # small sanity — last and first non-empty
    nz = [i for i, v in enumerate(series["expected_med"]) if v is not None]
    if nz:
        print(f"первое: {date_cols[nz[0]][1]} → expected {series['expected_med'][nz[0]]}%")
        print(f"последнее: {date_cols[nz[-1]][1]} → expected {series['expected_med'][nz[-1]]}%")
    return 0


if __name__ == "__main__":
    sys.exit(main())
