import json, collections, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
D = json.load(open('result.json')); A = json.load(open('ankety.json'))
C = D['canon']; idx = D['idx']; resp = set(D['resp'])
people = [p for p in D['people'] if p != 'Аноним']; N = len(people)
nm = lambda p: C[p][0]

doc = Document()
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
sec = doc.sections[0]; sec.left_margin = Cm(3); sec.right_margin = Cm(1.5); sec.top_margin = sec.bottom_margin = Cm(2)
def P(text='', bold=False, italic=False, align=None, size=None, indent=True, space=0):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(space); p.paragraph_format.line_spacing = 1.15
    if indent and align is None: p.paragraph_format.first_line_indent = Cm(1.25)
    if align == 'c': p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'r': p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align is None: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if not part: continue
        r = p.add_run(part.strip('*') if part.startswith('**') else part)
        r.bold = bold or part.startswith('**'); r.italic = italic
        if size: r.font.size = Pt(size)
    return p
def H(text): P(text, bold=True, align='c', space=6)
def table(rows, widths=None, size=11, header=True, bold_first_col=False):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ''
            p = c.paragraphs[0]; r = p.add_run(str(val)); r.font.size = Pt(size)
            if (header and i == 0) or (bold_first_col and j == 0): r.bold = True
            if widths: c.width = Cm(widths[j])
    doc.add_paragraph()
    return t

# ---------- титульный лист
for s in ['МИНИСТЕРСТВО ПРОСВЕЩЕНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ', 'ФЕДЕРАЛЬНОЕ ГОСУДАРСТВЕННОЕ БЮДЖЕТНОЕ ОБРАЗОВАТЕЛЬНОЕ УЧРЕЖДЕНИЕ ВЫСШЕГО ОБРАЗОВАНИЯ',
          '«ВОРОНЕЖСКИЙ ГОСУДАРСТВЕННЫЙ ПЕДАГОГИЧЕСКИЙ УНИВЕРСИТЕТ»', '', 'КАФЕДРА ОБЩЕЙ И ПЕДАГОГИЧЕСКОЙ ПСИХОЛОГИИ', '']:
    P(s, align='c', bold=bool(s))
P('ПРОИЗВОДСТВЕННАЯ ПЕДАГОГИЧЕСКАЯ ПРАКТИКА', bold=True, align='c'); P('', align='c')
P('Социометрическое изучение коллектива класса', bold=True, align='c', size=16); P('', align='c')
P('Направление подготовки: 44.03.05 Педагогическое образование (с двумя профилями подготовки)', indent=False)
P('Место прохождения практики: МБОУ «Образовательный центр им. А. Н. Косыгина», г. Красногорск, 7-Б класс', indent=False)
for _ in range(1): P('')
for s in ['Выполнила студентка заочной формы обучения,', '4 курс, 1 группа, гуманитарный факультет', 'Бычкова Ульяна Андреевна', '',
          'Руководитель по кафедре общей и педагогической психологии:', 'доцент кафедры общей и педагогической психологии', 'Колосова Е. В.', '', 'Оценка: ____________________', '', 'Подпись: ___________________']:
    p = P(s, indent=False); p.paragraph_format.left_indent = Cm(8); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
for _ in range(1): P('')
P('Воронеж – 2026', align='c')
doc.add_page_break()

# ---------- задание и описание
filled = len(A); absent = [p for p in people if p not in resp]
H('Задание по кафедре общей и педагогической психологии')
for s in ['1. Первичное ознакомление с классом и наблюдение за социально-психологическим климатом.',
          '2. Разработка анкеты из 12 вопросов, два из которых социометрические.',
          '3. Проведение анкетирования, обработка результатов, оформление социометрии.',
          '4. Разработка практических рекомендаций по работе с классом, подбор игр и упражнений на сплочение коллектива.']:
    P(s)
P(f'Анкетирование проведено 23 сентября 2026 года в 7-Б классе. Анкеты заполнили {filled} обучающихся '
  f'(одна анкета не подписана). В ответах на вопросы о выборе одноклассников названы ещё {len(absent)} обучающихся, '
  f'не заполнявших анкету: {", ".join(nm(p) for p in absent)}. Всего в исследование вошли {N} человек (Приложение 1).')
P('Анкета (Приложение 2) содержит 12 вопросов. Вопросы 11 и 12 — социометрические: вопрос 11 выявляет '
  'отношения по эмоциональному (неформальному) критерию, вопрос 12 — по деловому (учебному). Вопрос 5 («к кому '
  'из одноклассников ты можешь обратиться за помощью») использован как дополнительный. Каждый обучающийся мог '
  'сделать до трёх выборов по каждому социометрическому вопросу; учитывались только выборы одноклассников, которых '
  'удалось однозначно установить.')

# ---------- Приложение 1
H('Приложение 1. Список обучающихся, вошедших в исследование')
rows = [['№', 'Обучающийся', 'Пол', 'Анкета']]
for p in people: rows.append([idx[p], nm(p), C[p][1], '+' if p in resp else '—'])
rows.append(['А', 'Анкета без подписи', 'м (?)', '+'])
table(rows, [1.2, 7, 2, 2], size=11)
P('Знак «+» — анкета заполнена, «—» — анкету не заполнял, но назван одноклассниками. Знаком «(?)» отмечены '
  'фамилии, прочитанные с анкеты неуверенно: их нужно сверить со списком класса.', italic=True, indent=False, size=12)

# ---------- Приложение 2
H('Приложение 2. Анкета (образец бланка)')
Q = ['Анкета учащегося', 'Фамилия, имя: ______________________ Класс: ______ Дата: ______',
 '1. Нравится ли тебе учиться в этом классе? (да, очень; скорее да; скорее нет; нет)',
 '2. Как бы ты оценил(а) отношения с одноклассниками? (очень хорошие; хорошие; нормальные; не очень хорошие)',
 '3. Чувствуешь ли ты себя в классе комфортно? (да, всегда; обычно да; иногда нет; чаще нет)',
 '4. Бывают ли в классе ссоры и конфликты? (очень редко; иногда; часто; почти всегда)',
 '5. Если у тебя возникает проблема, к кому из одноклассников ты можешь обратиться за помощью? (можно указать несколько фамилий или написать «никто»)',
 '6. Как ты считаешь, есть ли в классе ученики, которых несправедливо обижают или не замечают? (нет; да, один-два человека; да, несколько человек; затрудняюсь ответить)',
 '7. Хотел(а) бы ты перейти в другой класс? (нет; скорее нет; скорее да; да)',
 '8. Как ты считаешь, умеют ли ребята в твоём классе работать вместе (например, делать общий проект)? (да, очень хорошо; нормально; не очень; совсем не умеют)',
 '9. Что бы ты хотел(а) изменить в своём классе? (можно написать несколько предложений)',
 '10. Как ты проводишь время на перемене? (общаюсь с друзьями из класса; общаюсь с друзьями из других классов; читаю / сижу в телефоне; другое)',
 '11. Если бы твой класс поехал на экскурсию в другой город, с кем из одноклассников ты хотел(а) бы ехать в одном купе (комнате)? Напиши 3 фамилии в порядке предпочтения.',
 '12. С кем из учеников вашего класса ты хотел(а) бы готовиться к контрольной работе или делать совместный проект? Напиши 3 фамилии в порядке предпочтения.']
for q in Q: P(q, indent=False, size=12)
P('Вопросы 11 и 12 — социометрические (эмоциональный и деловой критерии).', italic=True, indent=False, size=12)

# ---------- социоматрицы (альбомная ориентация)
def landscape():
    s = doc.add_section(); s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width, s.page_height = s.page_height, s.page_width
    s.left_margin = s.right_margin = Cm(1.2); s.top_margin = s.bottom_margin = Cm(1.5)
def portrait():
    s = doc.add_section(); s.orientation = WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = s.page_height, s.page_width
    s.left_margin = Cm(3); s.right_margin = Cm(1.5); s.top_margin = s.bottom_margin = Cm(2)
def matrix(q, title):
    r = D['R'][q]; M = r['M']; mutual = {tuple(x) for x in r['mutual']}
    P(title, bold=True, align='c', size=12)
    P('«+» — выбор; полужирным — взаимный выбор. Обучающиеся, не заполнявшие анкету, выборов не делали.', italic=True, align='c', size=9)
    choosers = [p for p in people if p in resp] + ['Аноним']
    rows = [['№', 'Кто выбирает'] + [str(idx[p]) for p in people] + ['Сделано']]
    for a in choosers:
        row = ['А' if a == 'Аноним' else idx[a], 'Без подписи' if a == 'Аноним' else nm(a)]
        for b in people:
            row.append('—' if a == b else ('+' if b in M.get(a, []) else ''))
        row.append(len(M.get(a, []))); rows.append(row)
    rows.append(['', 'Получено выборов'] + [r['got'].get(p, 0) for p in people] + [r['tot']])
    rows.append(['', 'Взаимных выборов'] + [sum(1 for m in mutual if p in m) for p in people] + [len(mutual)])
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'; t.autofit = False
    for j in range(len(rows[0])): t.columns[j].width = Cm(3.0) if j == 1 else Cm(0.62)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ''; pp = c.paragraphs[0]
            pp.paragraph_format.space_after = Pt(0); pp.paragraph_format.space_before = Pt(0); pp.paragraph_format.line_spacing = 1.0
            run = pp.add_run(str(v)); run.font.size = Pt(6.5)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
            if i == 0 or i >= len(rows) - 2: run.bold = True
            if 1 <= i < len(rows) - 2 and 2 <= j < len(row) - 1 and v == '+':
                a = choosers[i - 1]; b = people[j - 2]
                if tuple(sorted((a, b))) in mutual: run.bold = True; run.underline = True
            c.width = Cm(3.0) if j == 1 else Cm(0.62)
    doc.add_paragraph()
landscape()
matrix('q11', 'Приложение 3. Социоматрица выборов. Эмоциональный критерий (вопрос 11)')
doc.add_page_break()
matrix('q12', 'Приложение 4. Социоматрица выборов. Деловой критерий (вопрос 12)')
portrait()

# ---------- Приложение 5. Результаты
def stats(q):
    r = D['R'][q]; M = r['M']; tot = r['tot']; avg = tot / N
    to_resp = sum(1 for a in M for b in M[a] if b in resp)
    mf = sum(1 for a in M for b in M[a] if C[a][1] != C[b][1] and '?' not in (C[a][1], C[b][1]))
    return dict(tot=tot, made=r['made'], avg=avg, pairs=len(r['mutual']), kv=2 * len(r['mutual']) / tot,
                kv2=2 * len(r['mutual']) / to_resp, mf=mf, got=r['got'], mutual=r['mutual'])
S = {q: stats(q) for q in ('q11', 'q12')}
H('Приложение 5. Результаты обработки')
pct = lambda x: f'{round(x * 100)} %'
e, d = S['q11'], S['q12']
table([['Показатель', 'Эмоциональный критерий (в. 11)', 'Деловой критерий (в. 12)'],
       ['Обучающихся в исследовании', N, N], ['Заполнили анкету', filled, filled],
       ['Сделали выборы', e['made'], d['made']], ['Всего выборов', e['tot'], d['tot']],
       ['Среднее число полученных выборов', f"{e['avg']:.1f}".replace('.', ','), f"{d['avg']:.1f}".replace('.', ',')],
       ['Взаимных пар', e['pairs'], d['pairs']],
       ['Коэффициент взаимности (от всех выборов)', pct(e['kv']), pct(d['kv'])],
       ['То же, без учёта выборов в адрес не заполнявших анкету', pct(e['kv2']), pct(d['kv2'])],
       ['Выборы между мальчиками и девочками', f"{e['mf']} из {e['tot']} ({pct(e['mf'] / e['tot'])})", f"{d['mf']} из {d['tot']} ({pct(d['mf'] / d['tot'])})"]],
      [7, 4.5, 4.5], size=11)
def groups(s):
    g = s['got']; avg = s['avg']; out = {0: [], 1: [], 2: [], 3: []}
    for p in people:
        v = g.get(p, 0); k = 0 if v >= 2 * avg else 1 if v >= avg else 2 if v > 0 else 3
        out[k].append((v, nm(p)))
    def fmt(lst):
        by = collections.defaultdict(list)
        for v, n in lst: by[v].append(n)
        return '; '.join(f"{', '.join(sorted(ns))} ({v})" for v, ns in sorted(by.items(), reverse=True)) or '—'
    return {k: fmt(v) for k, v in out.items()}
ge, gd = groups(e), groups(d)
P('Статусные группы', bold=True, align='c')
table([['Группа', 'Эмоциональный критерий', 'Деловой критерий'],
       ['«Звёзды» (вдвое выше среднего)', ge[0], gd[0]], ['«Предпочитаемые» (на уровне среднего и выше)', ge[1], gd[1]],
       ['«Пренебрегаемые» (ниже среднего)', ge[2], gd[2]], ['«Изолированные» (0 выборов)', ge[3], gd[3]]], [4, 6, 6], size=10)
P('В скобках — число полученных выборов.', italic=True, indent=False, size=11)
P('Взаимные пары', bold=True, align='c')
for q, lab in (('q11', 'Эмоциональный критерий'), ('q12', 'Деловой критерий')):
    P(f"**{lab}:** " + '; '.join(f'{nm(a)} — {nm(b)}' for a, b in S[q]['mutual']) + '.')
P('Дополнительный вопрос 5 (к кому можно обратиться за помощью)', bold=True, align='c')
g5 = D['R']['q5']['got']
P('Чаще всего называли: ' + ', '.join(f'{nm(p)} ({v})' for p, v in sorted(g5.items(), key=lambda x: -x[1]) if v >= 3) +
  '. Ответ «никто» дали ' + str(sum(1 for a in A if any('никто' in x.lower() or 'никому' in x.lower() for x in a['q5']))) + ' обучающихся.')
P('Особенности заполнения', bold=True, align='c')
for s in [
  '— Одна анкета не подписана; её выборы внесены в социоматрицу отдельной строкой («Без подписи»).',
  '— Фамилии трёх обучающихся прочитаны неуверенно (Бакин Г., Тарариков Д., Хохлов М. — отмечены «(?)»); их нужно сверить со списком класса.',
  '— Не удалось установить, кого имели в виду: «Дима Черемисов» и «Дудак» (Хохлов М.), «Пансласт» (Нестеренко К.), «Саша …» (Бакин Г.). Эти выборы не засчитаны.',
  '— «Стёпа» (Хохлов М.) и «Карцев Стёпа» (Поливаев М.) засчитаны как выбор Карцева С.',
  '— Рягузов М. не ответил на вопросы 11 и 12; Белявцева В. и Поливаев М. не ответили на вопрос 12 («с кем угодно»); Тарариков Д. на вопрос 12 ответил «никто».',
  '— Мастеренко А. на вопросы 1, 2, 4, 7 и 8 отметила по два соседних варианта, каждый раз в сторону неблагополучия.']:
    P(s, indent=False, size=12)

# ---------- остальные вопросы
H('Приложение 6. Результаты по остальным вопросам анкеты')
def dist(q, opts):
    c = collections.Counter()
    for a in A:
        v = a[q]
        for o in opts:
            if o.lower() in v.lower(): c[o] += 1
    return c
QS = [('q1', '1. Нравится ли учиться в классе', ['Да, очень', 'Скорее да', 'Скорее нет', 'Нет']),
      ('q2', '2. Отношения с одноклассниками', ['Очень хорошие', 'Хорошие', 'Нормальные', 'Не очень хорошие']),
      ('q3', '3. Комфортно ли в классе', ['Да, всегда', 'Обычно да', 'Иногда нет', 'Чаще нет']),
      ('q4', '4. Ссоры и конфликты', ['Очень редко', 'Иногда', 'Часто', 'Почти всегда']),
      ('q6', '6. Есть ли те, кого обижают / не замечают', ['Нет', 'Да, один-два человека', 'Да, несколько человек', 'Затрудняюсь ответить']),
      ('q7', '7. Хотел(а) бы перейти в другой класс', ['Нет', 'Скорее нет', 'Скорее да', 'Да']),
      ('q8', '8. Умеют ли работать вместе', ['Да, очень хорошо', 'Нормально', 'Не очень', 'Совсем не умеют'])]
rows = [['Вопрос', 'Варианты ответа (число отметивших)']]
for q, lab, opts in QS:
    c = collections.Counter()
    for a in A:
        parts = [x.strip() for x in re.split(r' / ', re.sub(r'\s*\(.*?оба\)', '', a[q]))]
        for x in parts:
            for o in opts:
                if x.lower() == o.lower(): c[o] += 1
    rows.append([lab, '; '.join(f'{o.lower()} — {c[o]}' for o in opts)])
rows.append(['10. Как проводит перемену', 'с друзьями из класса — 16; с друзьями из других классов — 3; читаю / сижу в телефоне — 6; другое — 8 '
             '(«гуляю со Златой», «гуляю с подругами по школе», «хожу по школе», «хожу туда-сюда по коридору», «сижу в коридоре одна», «всё сразу» и др.)'])
table(rows, [5.5, 10.5], size=10)
P(f'Обучающиеся могли отметить несколько вариантов, поэтому сумма может превышать {filled}.', italic=True, indent=False, size=11)
P('Ответы на вопрос 9 («что бы ты хотел(а) изменить в классе»). Большинство ответили «ничего» или «всё хорошо». '
  'Содержательные ответы: «убрать пересадки» (2 человека), «чтобы некоторые девочки не общались группами», «хорошее отношение '
  'добавить, нормальных одноклассников, избавиться от буллинга», «чтобы [учителя] не придирались», «хотелось бы, чтобы были '
  'экскурсии», «чтобы КР и диктанты были реже, каникулы продлить», «чтобы была одна алгебра, геометрия».')

# ---------- социограммы
doc.add_page_break(); H('Социограммы')
for f in ('sociogramma_emocion.png', 'sociogramma_delovoy.png'):
    doc.add_picture(f, width=Cm(16)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
P('Номера соответствуют Приложению 1. Красные двусторонние стрелки — взаимные выборы, синие — односторонние. '
  'Пунктиром обведены обучающиеся, не заполнявшие анкету: они получали выборы, но сами их не делали. '
  'В скобках — число полученных выборов.', italic=True, indent=False, size=11)

# ---------- выводы
doc.add_page_break(); H('Выводы')
for s in open('vyvody.txt', encoding='utf-8').read().strip().split('\n\n'):
    if s.startswith('##'): H(s[2:].strip())
    else: P(s.strip())
doc.save('Sociometria_7B.docx'); print('ok')
